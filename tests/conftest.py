"""
Shared test tools: sample file makers and PDF checks.

The tests write real files and read real PDFs back on purpose: for a
converter the question is whether the content ends up in the PDF, not
whether a function was called.
"""

from __future__ import annotations

import datetime as dt
import sys
import zipfile
from pathlib import Path

import pytest

# The package is imported from the repository root (it is not installed).
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

# Accented probe text: "ő" and "ű" are exactly what reportlab's built-in
# fonts would turn into black squares.
HUN = "árvíztűrő tükörfúrógép ŐRÜLT"


@pytest.fixture
def workdir(tmp_path: Path) -> Path:
    return tmp_path


@pytest.fixture(scope="session")
def lo_profile(tmp_path_factory) -> Path:
    """A short LibreOffice profile path shared by the LibreOffice tests."""
    return tmp_path_factory.mktemp("lo")


@pytest.fixture(autouse=True)
def _libreoffice_mode(request, monkeypatch):
    """Hermetic tests: LibreOffice is hidden unless a test is marked
    `libreoffice` - then the real installation is used (or the test skips)."""
    from cornverter.converters import base

    base.reset_cancel()
    if request.node.get_closest_marker("libreoffice") is None:
        monkeypatch.setattr(base, "_SOFFICE_CACHE", "")
        return
    monkeypatch.setattr(base, "_SOFFICE_CACHE", None)
    if not base.find_soffice():
        pytest.skip("LibreOffice is not installed")
    monkeypatch.setenv("CORNVERTER_LO_PROFILE",
                       str(request.getfixturevalue("lo_profile")))


def pdf_text(pdf: Path) -> str:
    """The extracted text of every page of a PDF."""
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def pdf_pages(pdf: Path) -> int:
    from pypdf import PdfReader

    return len(PdfReader(str(pdf)).pages)


def assert_valid_pdf(pdf: Path) -> None:
    """An existing, non-empty, readable PDF."""
    assert pdf.exists(), f"not created: {pdf}"
    assert pdf.stat().st_size > 0, f"empty file: {pdf}"
    assert pdf.read_bytes()[:5] == b"%PDF-", f"no PDF header: {pdf}"
    assert pdf_pages(pdf) >= 1


def assert_no_missing_glyphs(pdf: Path) -> None:
    """A missing glyph shows up as a black square (U+25A0) - that is a bug."""
    assert "■" not in pdf_text(pdf), f"missing glyph in the PDF: {pdf}"


# ------------------------------------------------------------------ makers
def make_txt(path: Path, text: str = HUN) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def make_csv(path: Path, delimiter: str = ";") -> Path:
    rows = [
        ["Name", "Value", "Note"],
        [HUN, "42", "fine"],
        ["second", "7", "also fine"],
    ]
    path.write_text(
        "\n".join(delimiter.join(r) for r in rows) + "\n", encoding="utf-8")
    return path


def make_xlsx(path: Path, sheets: int = 2, with_empty: bool = True) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "First sheet"
    ws.append(["Name", "Date", "Number", "Flag"])
    ws.append([HUN, dt.date(2026, 8, 13), 3.0, True])
    ws.append(["decimal", dt.datetime(2026, 8, 13, 9, 30), 2.5, False])
    for i in range(2, sheets + 1):
        extra = wb.create_sheet(f"Sheet {i}")
        extra.append(["column", "value"])
        extra.append([f"row {i}", i])
    if with_empty:
        wb.create_sheet("Empty sheet")
    wb.save(path)
    return path


def make_png(path: Path, size=(640, 480), mode: str = "RGB") -> Path:
    from PIL import Image

    Image.new(mode, size, "white" if mode in ("RGB", "L") else None).save(path)
    return path


def make_multipage_tif(path: Path, pages: int = 3) -> Path:
    from PIL import Image

    frames = [Image.new("RGB", (600, 800), c)
              for c in ("white", "gray", "black")][:pages]
    frames[0].save(path, save_all=True, append_images=frames[1:])
    return path


_DOCX_XML = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
  <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
    <w:r><w:t>{heading}</w:t></w:r></w:p>
  <w:p><w:r><w:t>{body}</w:t></w:r><w:r><w:tab/><w:t>after a tab</w:t></w:r></w:p>
  <w:p/>
  <w:tbl>
    <w:tr><w:tc><w:p><w:r><w:t>Header A</w:t></w:r></w:p></w:tc>
           <w:tc><w:p><w:r><w:t>Header B</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:p><w:r><w:t>value 1</w:t></w:r></w:p></w:tc>
           <w:tc><w:p><w:r><w:t>value 2</w:t></w:r></w:p></w:tc></w:tr>
  </w:tbl>
</w:body></w:document>"""


def make_docx(path: Path, heading: str = "Heading", body: str = HUN) -> Path:
    """A minimal but genuinely structured .docx/.docm (no python-docx)."""
    xml = _DOCX_XML.format(heading=heading, body=body)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats'
            '.org/package/2006/content-types"><Default Extension="xml" '
            'ContentType="application/xml"/></Types>')
        z.writestr("word/document.xml", xml.encode("utf-8"))
    return path


_PPTX_SLIDE = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <p:cSld><p:spTree>
  <p:sp><p:txBody>
    <a:p><a:r><a:t>{title}</a:t></a:r></a:p>
    <a:p><a:r><a:t>{body}</a:t></a:r></a:p>
  </p:txBody></p:sp>
 </p:spTree></p:cSld></p:sld>"""

_PPTX_EMPTY_SLIDE = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <p:cSld><p:spTree/></p:cSld></p:sld>"""


def make_pptx(path: Path, slides: int = 2, empty: bool = False,
              note: str = "") -> Path:
    """A minimal but genuinely structured .pptx (no python-pptx)."""
    with zipfile.ZipFile(path, "w") as z:
        for i in range(1, slides + 1):
            xml = (_PPTX_EMPTY_SLIDE if empty else
                   _PPTX_SLIDE.format(title=f"Slide {i} title", body=HUN))
            z.writestr(f"ppt/slides/slide{i}.xml", xml.encode("utf-8"))
        if note:
            z.writestr(
                "ppt/notesSlides/notesSlide1.xml",
                _PPTX_SLIDE.format(title=note, body="").encode("utf-8"))
    return path


def make_rtf(path: Path, text: str = HUN) -> Path:
    """A realistic RTF with a Central European code page, escapes and tables.

    `\\fonttbl` and `\\stylesheet` deliberately hold SEVERAL `{...}` entries,
    like a real RTF saved by Word: a single-entry table once hid the bug of
    font names leaking into the finished PDF.
    """
    escaped = "".join(
        ch if ord(ch) < 128 else f"\\u{ord(ch)}?" for ch in text
    )
    path.write_text(
        "{\\rtf1\\ansi\\ansicpg1250\\deff0\\deflang1038\n"
        "{\\fonttbl"
        "{\\f0\\froman\\fprq2\\fcharset0 Times New Roman;}"
        "{\\f1\\fswiss\\fprq2\\fcharset0 Arial;}"
        "{\\f2\\fmodern\\fprq1\\fcharset0 Courier New;}"
        "{\\f3\\fnil\\fcharset2 Wingdings;}}\n"
        "{\\colortbl ;\\red0\\green0\\blue0;\\red255\\green0\\blue0;}\n"
        "{\\stylesheet{\\ql\\f0\\fs20 Normal;}{\\s1\\b\\fs28 heading 1;}"
        "{\\s2\\i\\fs24 heading 2;}}\n"
        "{\\*\\generator Riched20 10.0.19041;}\n"
        "{\\info{\\title Hidden title}{\\author Hidden author}}\n"
        "\\viewkind4\\uc1\\pard\\f0\\fs20 "
        f"{escaped}\\par\n"
        "Second paragraph\\tab after a tab.\\par\n"
        "}",
        encoding="cp1250", errors="replace")
    return path


def make_fake_xlsx_named_xls(path: Path) -> Path:
    """Real xlsx content with an .xls extension - a common mislabel."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mislabelled"
    ws.append(["Name", "Value"])
    ws.append([HUN, 11])
    wb.save(path)
    return path


def make_eml(path: Path, html: bool = False) -> Path:
    body = (f"<html><body><p>{HUN}</p></body></html>" if html
            else f"Plain text body: {HUN}")
    ctype = "text/html" if html else "text/plain"
    path.write_text(
        "From: sender@example.com\n"
        "To: recipient@example.com\n"
        "Subject: Test subject őrült\n"
        "Date: Thu, 13 Aug 2026 09:00:00 +0200\n"
        f"Content-Type: {ctype}; charset=utf-8\n"
        "MIME-Version: 1.0\n"
        "\n"
        f"{body}\n",
        encoding="utf-8")
    return path


def make_zip(path: Path, members: dict[str, bytes]) -> Path:
    with zipfile.ZipFile(path, "w") as z:
        for name, data in members.items():
            z.writestr(name, data)
    return path
