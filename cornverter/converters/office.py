"""
Office and ODF documents to PDF.

Handled: .xlsx .xls .xlsm .xlsb .ods .rtf .doc .docx .docm .odt .ppt .pptx
.pptm .odp

Two routes, in this order:

  1. **LibreOffice** (headless), when installed: formatting, column widths
     and images come through. .xlsx/.xlsm files are prepared with openpyxl
     first: landscape, fit to page width, and every column widened to its
     content - a column that is too narrow makes LibreOffice print "###" for
     dates and numbers and cut long text at the cell border. The PDF of a
     spreadsheet is still checked for "###"; if found, the built-in renderer
     is used instead, because silently lost content is worse than lost
     formatting.
  2. **Built-in fallback**, when LibreOffice is missing or fails: sheet.py
     (spreadsheets), wordml.py (Word), pptxml.py (presentations) and
     rtftext.py (RTF), all rendered with reportlab. The formatting is lost,
     the content is not, and the reason LibreOffice failed stays in the
     message.

Formats without a fallback (.doc, .ppt, .xlsb and the ODF formats) fail with
a clear "needs LibreOffice" message.
"""

from __future__ import annotations

import datetime as _dt
import importlib
import shutil
import tempfile
import time
from pathlib import Path

from .base import (
    ConversionResult,
    SofficeError,
    Status,
    find_soffice,
    long_path,
    run_soffice_convert,
)

# Spreadsheets whose page setup and column widths are prepared first.
_FIT_EXTS = {".xlsx", ".xlsm"}

# Spreadsheets whose LibreOffice PDF is checked for "###" (they have a
# built-in fallback to switch to).
_CHECK_CUT_CELLS = {".xlsx", ".xlsm", ".xls"}

# Extension -> built-in fallback module.
_FALLBACK = {
    ".xlsx": "sheet", ".xlsm": "sheet", ".xls": "sheet",
    ".docx": "wordml", ".docm": "wordml",
    ".pptx": "pptxml", ".pptm": "pptxml",
    ".rtf": "rtftext",
}

_NO_FALLBACK_HINT = (
    "This format needs LibreOffice (free: libreoffice.org) - install it and "
    "restart CornVerter."
)

# Rows sampled per sheet when estimating column widths.
_WIDTH_SAMPLE_ROWS = 500


def _cell_width(value) -> float:
    """Approximate printed width of a cell value, in characters."""
    if isinstance(value, bool):
        return 6
    if isinstance(value, _dt.datetime):
        return 19
    if isinstance(value, (_dt.date, _dt.time)):
        return 11
    if isinstance(value, int):
        return len(f"{value:,}") + 1
    if isinstance(value, float):
        return len(f"{value:,.2f}") + 1
    return max((len(line) for line in str(value).splitlines()), default=0)


def _widen_columns(ws) -> None:
    """Make every column at least as wide as its widest sampled value."""
    from openpyxl.utils import get_column_letter

    need: dict[int, float] = {}
    for row in ws.iter_rows(max_row=min(ws.max_row, _WIDTH_SAMPLE_ROWS)):
        for cell in row:
            if cell.value is None:
                continue
            need[cell.column] = max(need.get(cell.column, 0),
                                    _cell_width(cell.value))
    for col, width in need.items():
        dim = ws.column_dimensions[get_column_letter(col)]
        target = min(width * 1.2 + 2, 80)
        if (dim.width or 0) < target:
            dim.width = target


def _prepare_spreadsheet(src: Path) -> Path | None:
    """Fit-to-width page setup and content-wide columns on every sheet.

    Returns a temporary .xlsx path, or None (then the original is used).
    """
    try:
        import openpyxl
    except Exception:  # noqa: BLE001
        return None

    try:
        wb = openpyxl.load_workbook(long_path(src))
        for ws in wb.worksheets:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0  # no limit on the height
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_margins.left = 0.3
            ws.page_margins.right = 0.3
            ws.page_margins.top = 0.4
            ws.page_margins.bottom = 0.4
            _widen_columns(ws)
        tmp = Path(tempfile.mkdtemp(prefix="cv_xls_")) / (src.stem + ".xlsx")
        wb.save(tmp)
        return tmp
    except Exception:  # noqa: BLE001 - password, damage, macros: use the original
        return None


def _has_cut_cells(pdf: Path) -> bool:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(long_path(pdf)))
        return any("###" in (page.extract_text() or "")
                   for page in reader.pages[:100])
    except Exception:  # noqa: BLE001
        return False


def _fallback(path: Path, out_pdf: Path, reason: str | None) -> ConversionResult:
    """Use the built-in renderer for the format, if there is one."""
    prefix = f"{reason} " if reason else ""
    kind = _FALLBACK.get(path.suffix.lower())
    if kind is None:
        return ConversionResult(source=path, status=Status.FAILED,
                                message=f"{prefix}{_NO_FALLBACK_HINT}",
                                engine="libreoffice")

    module = importlib.import_module(f".{kind}", __package__)
    result = module.convert(path, out_pdf)
    if reason:
        if result.status is Status.FAILED:
            result.message = (f"{reason} The built-in renderer failed too: "
                              f"{result.message}")
        else:
            result.message = (f"{reason} Built-in renderer used: "
                              f"{result.message}").strip()
    return result


def _convert_with_soffice(path: Path, out_pdf: Path) -> ConversionResult:
    start = time.time()
    ext = path.suffix.lower()
    prepared = _prepare_spreadsheet(path) if ext in _FIT_EXTS else None

    tmp_out = Path(tempfile.mkdtemp(prefix="cv_office_"))
    try:
        produced = run_soffice_convert(prepared or path, tmp_out,
                                       convert_to="pdf")
        Path(long_path(out_pdf)).parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), str(long_path(out_pdf)))
        return ConversionResult(
            source=path, status=Status.SUCCESS, output=out_pdf,
            engine="libreoffice", duration_s=time.time() - start,
            message="All sheets in one PDF." if ext in _FIT_EXTS else "",
        )
    except SofficeError as exc:
        message = str(exc)
    except Exception as exc:  # noqa: BLE001
        message = f"{type(exc).__name__}: {exc}"
    finally:
        shutil.rmtree(tmp_out, ignore_errors=True)
        if prepared is not None:
            shutil.rmtree(prepared.parent, ignore_errors=True)
    return ConversionResult(source=path, status=Status.FAILED, message=message,
                            engine="libreoffice",
                            duration_s=time.time() - start)


def convert(path: Path, out_pdf: Path) -> ConversionResult:
    if not find_soffice():
        return _fallback(path, out_pdf, None)

    result = _convert_with_soffice(path, out_pdf)
    if result.status is Status.FAILED:
        # Damaged file, timeout, password... formatless content beats nothing.
        return _fallback(path, out_pdf, f"LibreOffice failed: {result.message}")
    if path.suffix.lower() in _CHECK_CUT_CELLS and _has_cut_cells(out_pdf):
        return _fallback(path, out_pdf,
                         "LibreOffice printed cut-off cells (###).")
    return result
